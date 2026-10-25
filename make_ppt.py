from pptx import Presentation
from pptx.util import Inches, Pt
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
def slide(title, bullets, sub=None):
    s = prs.slides.add_slide(BLANK)
    t = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1))
    t.text_frame.word_wrap = True; t.text_frame.text = title
    t.text_frame.paragraphs[0].font.size = Pt(36); t.text_frame.paragraphs[0].font.bold = True
    y = 1.6
    if sub:
        b = s.shapes.add_textbox(Inches(0.6), Inches(y), Inches(12), Inches(0.7)); b.text_frame.word_wrap=True; b.text_frame.text=sub; b.text_frame.paragraphs[0].font.size=Pt(20); y += 0.9
    for line in bullets:
        b = s.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.8), Inches(0.8)); b.text_frame.word_wrap=True; b.text_frame.text="•  "+line; b.text_frame.paragraphs[0].font.size=Pt(20); y += 0.85
    return s
slide("GradeGauge", ["Student Exam Score Predictor"], "An AI/ML project by Vikranth Reddy — B.Tech CSE (AI & ML), Woxsen University")
slide("The Problem", ["Students don't know where they stand before exams", "Study effort is guesswork — which habit actually moves the score?", "Teachers spot at-risk students too late to intervene"])
slide("The Solution", ["GradeGauge predicts the exam score (0–100) from 5 study inputs", "Interactive sliders: change a habit, see the predicted score instantly", "Runs fully in the browser — no backend, no data leaves the page"])
slide("Tech Stack", ["Python + scikit-learn: ridge regression (alpha = 1.0)", "Synthetic dataset: 500 student records (seed 42)", "HTML + CSS + vanilla JavaScript demo", "GitHub + Vercel hosting"])
slide("Architecture", ["Inputs: study hours, attendance, sleep, previous score, assignments", "Standardization (train mean / std) → ridge regression → clip to 0–100", "Model exported as model.json; the browser demo reproduces it exactly"])
slide("Demo", ["Try the presets: low-effort, average, and top student", "Average student (6 hrs, 85% attendance) → 96.4 predicted", "Top student preset → 100 (capped) — the model rewards consistent habits"])
slide("Results", ["5-fold cross-validated R² = 0.871", "Mean absolute error = 4.54 marks (baseline: 13.94 — about 3× better)", "RMSE = 5.81 marks", "Strongest drivers: study hours and previous score"])
slide("Future Scope", ["Train on real (anonymized) student data with consent", "Add subject-wise predictions and revision planners", "Explain each prediction: which habit to improve first", "Early-warning alerts for at-risk students"])
slide("Thank You", ["GradeGauge — predict, understand, improve."], "Vikranth Reddy · github.com/HVRGAMING2705/gradegauge")
out = "/home/hatch/workspace/goals/daily-ai-ml-project-factory/files/AIML_GradeGauge_2026-10-25.pptx"
prs.save(out); print("saved", out, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
